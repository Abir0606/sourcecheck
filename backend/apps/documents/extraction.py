"""Safe file validation + text extraction + chunking.

Extraction is best-effort and never raises raw library errors to callers; it
returns structured results so the pipeline can set a clear document status.
"""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass

from django.conf import settings

logger = logging.getLogger("apps.documents.extraction")


class ExtractionError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class ExtractionResult:
    pages: list[tuple[int | None, str]]
    ocr_used: bool = False


def _sniff_content_type(fileobj, filename: str) -> str:
    """Magic-byte check when libmagic is present; otherwise trust the extension."""
    try:
        import magic  # type: ignore

        head = fileobj.read(4096)
        fileobj.seek(0)
        return magic.from_buffer(head, mime=True)
    except Exception:
        fileobj.seek(0)
        return ""


def validate_upload(*, fileobj, filename: str, declared_type: str = "") -> None:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext not in settings.ALLOWED_UPLOAD_EXTENSIONS:
        raise ExtractionError(
            "unsupported_type",
            f"Unsupported file type '{ext or filename}'. Allowed: "
            + ", ".join(settings.ALLOWED_UPLOAD_EXTENSIONS),
        )
    fileobj.seek(0, io.SEEK_END)
    size = fileobj.tell()
    fileobj.seek(0)
    if size == 0:
        raise ExtractionError("empty_file", "The uploaded file is empty.")
    if size > settings.MAX_UPLOAD_BYTES:
        raise ExtractionError(
            "too_large", f"File exceeds the {settings.MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit."
        )


def extract(fileobj, *, filename: str, content_type: str, ocr) -> ExtractionResult:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    fileobj.seek(0)
    if ext == ".pdf":
        return _extract_pdf(fileobj, ocr)
    if ext == ".docx":
        return _extract_docx(fileobj)
    if ext == ".xlsx":
        return _extract_xlsx(fileobj)
    if ext in (".txt", ".md", ".csv", ".rtf", ".odt"):
        data = fileobj.read()
        text = data.decode("utf-8", errors="replace")
        return ExtractionResult(pages=[(None, text)])
    raise ExtractionError("unsupported_type", f"Unsupported file type '{ext}'.")


def _extract_pdf(fileobj, ocr) -> ExtractionResult:
    from pypdf import PdfReader

    try:
        reader = PdfReader(fileobj)
    except Exception as exc:
        raise ExtractionError("corrupt_file", "The PDF could not be read.") from exc

    pages: list[tuple[int | None, str]] = []
    for idx, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append((idx, text))

    total = sum(len(t.strip()) for _, t in pages)
    # Scanned PDFs yield little/no native text -> try OCR when available.
    if total < 20 and getattr(ocr, "available", False):
        fileobj.seek(0)
        try:
            ocr_pages = ocr.extract(fileobj, content_type="application/pdf")
            if ocr_pages:
                return ExtractionResult(pages=ocr_pages, ocr_used=True)
        except Exception:  # pragma: no cover - depends on tesseract/poppler
            logger.warning("OCR failed", exc_info=True)
    return ExtractionResult(pages=pages)


def _extract_docx(fileobj) -> ExtractionResult:
    import docx

    try:
        document = docx.Document(fileobj)
    except Exception as exc:
        raise ExtractionError("corrupt_file", "The DOCX file could not be read.") from exc
    paras = [p.text for p in document.paragraphs]
    return ExtractionResult(pages=[(None, "\n".join(paras))])


def _extract_xlsx(fileobj) -> ExtractionResult:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError("corrupt_file", "The XLSX file could not be read.") from exc
    pages: list[tuple[int | None, str]] = []
    for sheet in wb.worksheets:
        lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append("\t".join(cells))
        pages.append((None, "\n".join(lines)))
    return ExtractionResult(pages=pages)


def chunk_pages(
    pages: list[tuple[int | None, str]],
    *,
    target_chars: int | None = None,
    overlap_chars: int | None = None,
) -> list[dict]:
    """Split page text into overlapping passages, preserving page + char offsets."""
    target_chars = target_chars or settings.PASSAGE_TARGET_CHARS
    overlap_chars = overlap_chars or settings.PASSAGE_OVERLAP_CHARS
    passages: list[dict] = []
    global_index = 0
    for page_number, page_text in pages:
        text = page_text or ""
        if not text.strip():
            continue
        start = 0
        length = len(text)
        while start < length:
            end = min(start + target_chars, length)
            if end < length:
                window = text.rfind(" ", start + target_chars // 2, end)
                if window > start:
                    end = window
            chunk = text[start:end].strip()
            if chunk:
                passages.append(
                    {
                        "page_number": page_number,
                        "chunk_index": global_index,
                        "text": chunk,
                        "start_char": start,
                        "end_char": end,
                        "token_count": max(1, len(chunk.split())),
                    }
                )
                global_index += 1
            if end >= length:
                break
            start = max(end - overlap_chars, start + 1)
    return passages
